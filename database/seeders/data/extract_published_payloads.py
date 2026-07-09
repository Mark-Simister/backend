"""
Extracts the Published_Review_Payloads sheet from the BeastieRated Production OS
spreadsheet into published_review_payloads.json, consumed by
PublishedReviewPayloadSeeder. Each populated row becomes one public review
page (served at /review/{review_slug}).

The pipeline pre-builds review_page_json + schema_json_ld etc. in the sheet;
this script parses those JSON columns (repairing a known tail-corruption in
review_page_json) so the payloads land in the DB as clean, decodable JSON.

Run:  python backend/database/seeders/data/extract_published_payloads.py
"""
import openpyxl, json, os, glob

HERE = os.path.dirname(__file__)
SHEET_DIR = os.path.abspath(os.path.join(HERE, "..", "..", "..", "public", "product_review"))
OUT = os.path.join(HERE, "published_review_payloads.json")
SHEET = "Published_Review_Payloads"

JSON_COLS = {
    "qa_errors_json", "qa_warnings_json", "review_page_json", "schema_json_ld",
    "commerce_json", "sources_json", "safety_json", "disclosure_json",
}
INT_COLS = {"public_rating_count", "analysed_evidence_count", "source_count"}
FLOAT_COLS = {"final_beastie_score", "public_score"}


def find_xlsx():
    cands = glob.glob(os.path.join(SHEET_DIR, "BeastieRated_Production_OS*.xlsx"))
    return sorted(cands, key=os.path.getmtime, reverse=True)[0] if cands else None


def parse_json(col, s):
    """Decode a JSON cell, repairing the known review_page_json tail corruption
    (`"qa_warnings".[]"]}}` where a ':' became '.' plus a stray '"]')."""
    if s in (None, ""):
        return None, None
    if not isinstance(s, str):
        return s, None
    try:
        return json.loads(s), None
    except json.JSONDecodeError as e:
        repaired = s.replace('"qa_warnings".[]"]}}', '"qa_warnings":[]}}')
        try:
            return json.loads(repaired), f"{col}: repaired tail corruption"
        except json.JSONDecodeError:
            return None, f"{col}: UNPARSEABLE ({e}) — stored null"


def coerce(col, v):
    if v in (None, ""):
        return None
    if col in INT_COLS:
        try:
            return int(round(float(v)))
        except (TypeError, ValueError):
            return None
    if col in FLOAT_COLS:
        try:
            return round(float(v), 2)
        except (TypeError, ValueError):
            return None
    return v


def main():
    xlsx = find_xlsx()
    if not xlsx:
        print("No BeastieRated_Production_OS*.xlsx found in", SHEET_DIR)
        return
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        print(f"Sheet '{SHEET}' not found in {os.path.basename(xlsx)}")
        return
    ws = wb[SHEET]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]

    out, warnings = [], []
    for row in rows:
        d = dict(zip(header, row))
        slug = d.get("review_slug")
        pid = d.get("published_review_id")
        if not slug or not pid:
            continue  # empty template row
        rec = {}
        for col in header:
            if not col:
                continue
            v = d.get(col)
            if col in JSON_COLS:
                parsed, warn = parse_json(col, v)
                rec[col] = parsed
                if warn:
                    warnings.append(f"{slug[:40]}… {warn}")
            else:
                rec[col] = coerce(col, v)
        out.append(rec)

    json.dump(out, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"source: {os.path.basename(xlsx)}")
    print(f"payloads extracted: {len(out)}  ->  {OUT}")
    for w in warnings:
        print("  warn:", w)
    if out:
        r = out[0]
        print(f"  e.g. {r.get('character_name')} · {r.get('review_slug')[:50]}… · score {r.get('final_beastie_score')}")


if __name__ == "__main__":
    main()
